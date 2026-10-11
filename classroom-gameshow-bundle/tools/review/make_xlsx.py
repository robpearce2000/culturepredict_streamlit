#!/usr/bin/env python3
"""Builds review/science-review.xlsx from the per-pack review files in review/science/*.json.
   python3 tools/review/make_xlsx.py"""
import glob, json, subprocess, collections
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

SPECS = ['AQA 8461 Biology, version 1.0 (21 April 2016)', 'AQA 8462 Chemistry, version 1.1 (4 October 2019)',
         'AQA 8463 Physics, version 1.1 (30 September 2019)', 'AQA 8464 Combined Science: Trilogy, version 1.1 (4 October 2019)',
         'Edexcel 1BI0 Biology, 1CH0 Chemistry and 1PH0 Physics, Issue 4', 'Edexcel 1SC0 Combined Science, Issue 6']
ORDER = {'Wrong': 0, 'Doubtful': 1, 'Minor': 2}
HDR = ['ID', 'Board', 'Subject', 'Course', 'Tier', 'Topic', 'Spec ref', 'Question', 'Answer', 'Accepted answers', 'Grade', 'Reason', 'Suggested fix', "Rob's decision", "Rob's note"]
W = [22, 9, 10, 15, 11, 18, 11, 48, 24, 28, 10, 46, 46, 18, 30]
f = Font(name='Arial', size=10); fb = Font(name='Arial', size=10, bold=True, color='FFFFFF')
hfill = PatternFill('solid', fgColor='1F3864'); ifill = PatternFill('solid', fgColor='FFF2CC')
gfill = {'Wrong': 'F4CCCC', 'Doubtful': 'FCE5CD', 'Minor': 'D9EAD3'}
thin = Side(style='thin', color='BFBFBF'); bd = Border(top=thin, bottom=thin, left=thin, right=thin)
wrap = Alignment(wrap_text=True, vertical='top')

rows = collections.defaultdict(list); totals = collections.Counter(); reviewed = collections.Counter(); calcs = collections.Counter()
for path in sorted(glob.glob('review/science/*.json')):
    rv = json.load(open(path))
    data = json.loads(subprocess.check_output(['node', 'tools/review/export.js', rv['pack']]))
    subj = rv['pack'].split('/')[2].capitalize(); board = 'AQA' if '/aqa/' in rv['pack'] else 'Edexcel'
    totals[(subj, board)] += data['count']; reviewed[(subj, board)] += rv.get('reviewed', 0); calcs[(subj, board)] += len(rv.get('calcs', {}))
    qs = {q['id']: q for q in data['questions']}
    for fl in rv.get('flags', []):
        q = qs[fl['id']]
        tier = q['tier'] if q['course'] == 'Separate only' or q['tier'] == q['combTier'] else f"Separate: {q['tier']}; Combined: {q['combTier']}"
        rows[subj].append([q['id'], board, subj, q['course'], tier, data['topic'], q['ref'] + (f" (Combined {q['combRef']})" if q['combRef'] and q['combRef'] != q['ref'] else ''),
                           q['q'], q['a'], '; '.join(q['accept']), fl['grade'], fl['reason'], fl['fix'], None, None])
wb = Workbook(); summ = wb.active; summ.title = 'Summary'
for subj in ('Biology', 'Chemistry', 'Physics'):
    ws = wb.create_sheet(subj); ws.append(HDR)
    for c, w in zip('ABCDEFGHIJKLMNO', W): ws.column_dimensions[c].width = w
    for cell in ws[1]: cell.font = fb; cell.fill = hfill; cell.alignment = wrap; cell.border = bd
    ws.freeze_panes = 'B2'
    dv = DataValidation(type='list', formula1='"Use suggested fix,Keep as is,Remove,I\'ll edit"', allow_blank=True); ws.add_data_validation(dv)
    for r in sorted(rows[subj], key=lambda r: (ORDER[r[10]], r[1], r[0])):
        ws.append(r); i = ws.max_row
        for cell in ws[i]: cell.font = f; cell.alignment = wrap; cell.border = bd
        ws.cell(i, 11).fill = PatternFill('solid', fgColor=gfill[r[10]]); ws.cell(i, 14).fill = ifill; ws.cell(i, 15).fill = ifill
        dv.add(ws.cell(i, 14))
    ws.auto_filter.ref = f'A1:O{max(ws.max_row, 2)}'
summ.append(['Science question review (Biology, Chemistry, Physics; AQA and Edexcel; separate and Combined Science)'])
summ['A1'].font = Font(name='Arial', size=14, bold=True)
notes = ['How to fill in: on each subject sheet, choose a value in the yellow "Rob\'s decision" column: Use suggested fix, Keep as is, Remove, or I\'ll edit (put your wording in "Rob\'s note"). Then upload the file back to review/ with the same name.',
         'Method: every question was reviewed by an independent reviewer, pack by pack, against the specification text; every numerical answer (found by its content, not the Calc tag) was recalculated with a script (tools/review/calc_check.py). Only flagged questions are listed, Wrong first, then Doubtful, then Minor.',
         'Specifications used: ' + '; '.join(SPECS) + '.',
         'Question, reviewed and recalculated counts are counted from the packs and review files (hardcoded); flag counts are formulas over the subject sheets.']
for n_ in notes:
    summ.append([n_]); r = summ.max_row; summ.cell(r, 1).font = f; summ.cell(r, 1).alignment = wrap
    summ.merge_cells(start_row=r, start_column=1, end_row=r, end_column=9); summ.row_dimensions[r].height = 45
summ.append([])
summ.append(['Subject', 'Board', 'Questions', 'Reviewed', 'Recalculated', 'Wrong', 'Doubtful', 'Minor', 'All flags'])
hr = summ.max_row
for cell in summ[hr]: cell.font = fb; cell.fill = hfill; cell.border = bd; cell.alignment = wrap
first = hr + 1
for subj in ('Biology', 'Chemistry', 'Physics'):
    for b in ('AQA', 'Edexcel'):
        r = summ.max_row + 1
        summ.append([subj, b, totals[(subj, b)], reviewed[(subj, b)], calcs[(subj, b)]] +
                    [f'=COUNTIFS({subj}!$B$2:$B$2000,$B{r},{subj}!$K$2:$K$2000,"{g}")' for g in ('Wrong', 'Doubtful', 'Minor')] + [f'=SUM(F{r}:H{r})'])
last = summ.max_row
summ.append(['Total', None] + [f'=SUM({c}{first}:{c}{last})' for c in 'CDEFGHI'])
for row in summ.iter_rows(min_row=first, max_row=summ.max_row):
    for cell in row: cell.font = f; cell.border = bd
for cell in summ[summ.max_row]: cell.font = Font(name='Arial', size=10, bold=True)
for c, w in zip('ABCDEFGHI', [12, 10, 11, 11, 13, 9, 10, 9, 10]): summ.column_dimensions[c].width = w
wb.calculation.fullCalcOnLoad = True
wb.save('review/science-review.xlsx')
print({k: len(v) for k, v in rows.items()}, dict(totals), dict(reviewed))
